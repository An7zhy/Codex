from dataclasses import replace
from io import BytesIO
import json
import math
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.analysis import analyze_shipments, calculate_kpis
from src.decisions import ACTIONS
from src.demo import DEMO_AS_OF, demo_xlsx_bytes, generate_demo_data
from src.history import CaseInput, HistoryStore
from src.ingestion import load_shipments
from src.priority import COMPONENTS, PriorityConfig, prioritize_shipments
from src.schema import DataInputError

ROOT = Path(__file__).resolve().parents[1]


def test_extra_csv_fields_cannot_silently_shift_shipment_columns():
    lines = generate_demo_data().head(1).to_csv(index=False).strip().splitlines()
    lines[1] += ",unexpected-extra-field"
    with pytest.raises(DataInputError):
        load_shipments(("\n".join(lines) + "\n").encode(), "ragged.csv")


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_duplicate_named_headers_are_rejected_before_pandas_mangling(extension):
    data = generate_demo_data().head(1)
    if extension == "csv":
        lines = data.to_csv(index=False).strip().splitlines()
        lines[0] += ",shipping_cost"
        lines[1] += ",999999"
        content = ("\n".join(lines) + "\n").encode()
    else:
        book = load_workbook(BytesIO(demo_xlsx_bytes(data)))
        sheet = book.worksheets[0]
        column = sheet.max_column + 1
        sheet.cell(1, column, "shipping_cost")
        sheet.cell(2, column, 999999)
        output = BytesIO()
        book.save(output)
        content = output.getvalue()
    with pytest.raises(DataInputError, match="unique"):
        load_shipments(content, "ambiguous." + extension)


def test_damaged_excel_xml_is_an_input_error_not_an_uncaught_exception():
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(demo_xlsx_bytes())) as original, zipfile.ZipFile(output, "w") as broken:
        for name in original.namelist():
            broken.writestr(name, b"<broken" if name == "xl/workbook.xml" else original.read(name))
    with pytest.raises(DataInputError):
        load_shipments(output.getvalue(), "damaged.xlsx")


def test_extreme_costs_cannot_overflow_kpis(make_rows):
    raw = make_rows(*[{"shipping_cost": 1e308}] * 8, {"shipping_cost": 100})
    result = analyze_shipments(raw, DEMO_AS_OF)
    kpis = calculate_kpis(result.shipments)
    assert math.isfinite(kpis["total_shipping_cost"])
    assert kpis["total_shipping_cost"] == 100
    assert kpis["cost_eligible_shipments"] == 1
    assert len(result.issues.loc[result.issues["code"].eq("invalid_cost")]) == 8


def test_unchanged_repeated_save_click_reuses_case(tmp_path, monkeypatch):
    database = tmp_path / "double-click.sqlite3"
    monkeypatch.setenv("LOGISTICS_DB_PATH", str(database))
    app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=30)
    for _ in range(2):
        next(button for button in app.button if button.label == "Save decision to local history").click().run(timeout=30)
        assert not app.exception
    assert len(HistoryStore(database).list_cases()) == 1
    next(button for button in app.button if button.label == "Apply priority settings").click().run(timeout=30)
    next(button for button in app.button if button.label == "Save decision to local history").click().run(timeout=30)
    assert len(HistoryStore(database).list_cases()) == 1
    app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=30)
    next(button for button in app.button if button.label == "Save decision to local history").click().run(timeout=30)
    assert not app.exception
    assert len(HistoryStore(database).list_cases()) == 1
    notes = next(widget for widget in app.text_area if widget.label == "Decision notes")
    notes.set_value("Changed decision rationale creates a distinct historical decision.")
    next(button for button in app.button if button.label == "Save decision to local history").click().run(timeout=30)
    assert not app.exception
    assert len(HistoryStore(database).list_cases()) == 2


@pytest.mark.parametrize("field", ["assumptions_json", "priority_snapshot_json"])
def test_invalid_saved_metadata_is_rejected(tmp_path, field):
    store = HistoryStore(tmp_path / "metadata.sqlite3")
    with pytest.raises(ValueError):
        store.create_case(CaseInput("S1", "Overdue in transit", ACTIONS[0], **{field: "not-json"}))
    assert store.list_cases().empty


def test_mutated_priority_config_is_revalidated(make_rows):
    result = analyze_shipments(make_rows({"status": "In Transit", "actual_delivery": ""}), DEMO_AS_OF)
    config = PriorityConfig()
    config.weights.update(dict.fromkeys(COMPONENTS, 0))
    with pytest.raises(ValueError):
        prioritize_shipments(result.shipments, result.cutoff, config)


@pytest.mark.parametrize("component,expected", [("urgency", 100), ("delay", 50), ("financial", 50), ("value", 50), ("downstream", 40)])
def test_independent_component_weights_have_hand_calculated_results(make_rows, component, expected):
    raw = make_rows({"status": "In Transit", "actual_delivery": "", "delivery_deadline": "2026-10-02T00:00:00Z",
                     "estimated_delay_hours": 36, "estimated_financial_impact": 5000,
                     "financial_impact_basis": "Synthetic penalty assumption", "shipment_value": 50000,
                     "downstream_dependency_count": 2, "downstream_risk_basis": "Two synthetic dependencies"})
    result = analyze_shipments(raw, DEMO_AS_OF)
    config = replace(PriorityConfig(), weights={key: 100 if key == component else 0 for key in COMPONENTS})
    assert prioritize_shipments(result.shipments, result.cutoff, config).iloc[0]["priority_score"] == expected


def test_sqlite_case_survives_separate_processes(tmp_path):
    database = tmp_path / "process-restart.sqlite3"
    creator = """import sys,json
from src.history import HistoryStore,CaseInput
from src.decisions import ACTIONS
s=HistoryStore(sys.argv[1])
print(s.create_case(CaseInput('PROCESS-CASE','Overdue in transit',ACTIONS[0],estimated_cost=10)))
"""
    first = subprocess.run([sys.executable, "-c", creator, str(database)], cwd=ROOT, capture_output=True, text=True, check=True)
    case_id = first.stdout.strip()
    updater = """import sys
from src.history import HistoryStore
s=HistoryStore(sys.argv[1]);s.update_case(sys.argv[2],'Resolved',12,24,'Synthetic persisted outcome',expected_version=1)
"""
    subprocess.run([sys.executable, "-c", updater, str(database), case_id], cwd=ROOT, check=True)
    reader = """import sys,json
from src.history import HistoryStore
print(json.dumps(HistoryStore(sys.argv[1]).get_case(sys.argv[2])))
"""
    last = subprocess.run([sys.executable, "-c", reader, str(database), case_id], cwd=ROOT, capture_output=True, text=True, check=True)
    record = json.loads(last.stdout)
    assert record["resolution_status"] == "Resolved"
    assert record["actual_cost"] == 12
    assert record["actual_delay"] == 24
    assert record["estimated_cost"] == 10
