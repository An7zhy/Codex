from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.analysis import analyze_shipments, calculate_kpis, exception_shipments
from src.demo import DEMO_AS_OF, demo_xlsx_bytes, generate_demo_data
from src.exports import csv_bytes
from src.ingestion import load_shipments
from src.reporting import RuleBasedReportGenerator
from src.schema import DataInputError, REQUIRED_COLUMNS


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_demo_upload_to_metrics_exceptions_and_report(extension):
    demo = generate_demo_data()
    content = demo.to_csv(index=False).encode() if extension == "csv" else demo_xlsx_bytes(demo)
    loaded = load_shipments(content, "synthetic_shipments." + extension)
    result = analyze_shipments(loaded, DEMO_AS_OF)
    k = calculate_kpis(result.shipments)
    assert len(loaded) == 160
    assert k["total_shipments"] == 158
    assert k["delivered_shipments"] == 91
    assert k["delayed_shipments"] == 47
    assert k["overdue_shipments"] == 30
    assert k["cost_anomalies"] == 4
    assert k["total_shipping_cost"] == pytest.approx(539830.75)
    assert not exception_shipments(result.shipments).empty
    report = RuleBasedReportGenerator().generate(result.shipments, result.issues, result.as_of, "Synthetic Demo Data")
    for section in ["Executive Summary", "Operational KPIs", "Key Exceptions", "Carrier Performance", "Data Quality Issues", "Recommended Next Actions"]:
        assert "## " + section in report
    assert "| Total Shipments | 158 |" in report
    assert "539,830.75" in report
    assert "Synthetic Demo Data" in report
    processed = pd.read_csv(BytesIO(csv_bytes(result.shipments)))
    assert len(processed) == 160
    assert "shipping_cost_raw" in processed


def test_demo_is_reproducible_and_excel_is_labelled():
    pd.testing.assert_frame_equal(generate_demo_data(), generate_demo_data())
    workbook = load_workbook(BytesIO(demo_xlsx_bytes()))
    assert workbook.sheetnames[0] == "Shipments"
    assert workbook["Shipments"].max_row == 161
    assert "Synthetic Demo Data" in workbook["Read me"]["B2"].value
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)


@pytest.mark.parametrize("content,name", [(b"", "empty.csv"), (b"   \n", "empty.csv"),
                                        (b"not-an-excel-file", "bad.xlsx"),
                                        (b"x\n1", "wrong.csv"),
                                        (b"hello", "wrong.txt"), (b"\xff\xfe", "bad.csv")])
def test_unreadable_or_wrong_schema_inputs_have_actionable_errors(content, name):
    with pytest.raises(DataInputError):
        load_shipments(content, name)


def test_header_only_csv_is_valid_empty_data():
    content = (",".join(REQUIRED_COLUMNS) + "\n").encode()
    loaded = load_shipments(content, "headers.csv")
    result = analyze_shipments(loaded, DEMO_AS_OF)
    assert calculate_kpis(result.shipments)["total_shipments"] == 0
    report = RuleBasedReportGenerator().generate(result.shipments, result.issues, result.as_of)
    assert "| On-Time Delivery Rate | N/A |" in report
    assert "| Total Shipments | 0 |" in report


def test_csv_preserves_leading_zero_ids_and_normalizes_headers(make_rows):
    data = make_rows({"shipment_id": "000001"})
    data.columns = [" " + column.upper() + " " for column in data]
    loaded = load_shipments(data.to_csv(index=False).encode("utf-8-sig"), "upload.CSV")
    assert loaded.iloc[0]["shipment_id"] == "000001"


def test_csv_formula_text_is_escaped_but_numbers_remain_numeric():
    data = pd.DataFrame({"text": ["=HYPERLINK(1)", " @cmd", "+calc", "normal", "\t=2+2"], "number": [-5, 2, 3, 4, 5]})
    loaded = pd.read_csv(BytesIO(csv_bytes(data)))
    assert loaded["text"].iloc[0].startswith("'=")
    assert loaded["text"].iloc[1].startswith("'")
    assert loaded["text"].iloc[2].startswith("'+")
    assert loaded["text"].iloc[4].startswith("'")
    assert loaded["number"].iloc[0] == -5
    assert loaded["text"].iloc[3] == "normal"


def test_filtered_report_uses_only_selected_rows_and_escapes_user_content(make_rows):
    result = analyze_shipments(make_rows({"carrier": "<script>alert(1)</script>"},
                                        {"shipment_id": "EXCLUDE", "actual_delivery": ""}), DEMO_AS_OF)
    selected = result.shipments.iloc[[0]]
    report = RuleBasedReportGenerator().generate(selected, result.issues, result.as_of, "Upload", "Filtered")
    assert "| Total Shipments | 1 |" in report
    assert "<script>" not in report
    assert "&lt;script&gt;" in report
    assert "missing\\_actual\\_delivery" not in report
