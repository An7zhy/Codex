"""Strict uploaded planning contract. No FX/unit/default-timezone inference.

All timestamps are offset-aware ISO instants. A service is a shared resource
with one expected transit duration; order-specific quotes map to Candidates.
Verification flags are declarations by the uploader, never third-party proof.
"""

import csv
from dataclasses import asdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import re

import pandas as pd

from src.ingestion import MAX_FILE_BYTES, _validate_named_headers, validate_xlsx_archive
from .models import Candidate, Constraints, Order, PlanningData, Preferences, Service

SCHEMA_VERSION = "uploaded-planning-v1"
HEADERS = {
    "shipments": ["shipment_id", "origin", "destination", "weight", "weight_unit", "volume", "volume_unit", "delivery_deadline"],
    "transport_options": ["shipment_id", "service_id", "quoted_cost", "currency", "expected_transit_hours", "time_unit", "quote_status", "quote_reference", "time_status", "time_reference"],
    "capacity": ["service_id", "origin", "destination", "transport_mode", "weight_capacity", "weight_unit", "volume_capacity", "volume_unit", "shipment_capacity", "capacity_status", "capacity_reference"],
    "constraints": ["planning_start", "budget", "currency", "allow_unassigned", "unassigned_penalty", "penalty_currency", "cost_weight", "time_weight", "time_limit_seconds", "source_label", "data_kind"],
}
LIMITS = {"shipments": 500, "transport_options": 5000, "capacity": 100, "constraints": 1}
POLICY = [
    "USD / kg / m3 / h only. No FX or unit inference/conversion. Decimal point numbers only.",
    "Offset-aware ISO timestamps required; date-only/naive dates rejected. Every order starts at the common planning_start.",
    "Shared service IDs share weight, volume and slots, with one common expected transit duration. Capacity is not repeated per order.",
    "Capacity without any quoted option is retained in source records but omitted from the model; utilization denominators cover modeled services only.",
    "Expected-time deadlines are hard eligibility constraints, not arrival guarantees. Demands/costs round up, capacities/budget/deadlines down to the existing solver units.",
    "estimate and user_verified are uploader declarations with references, not independent verification by this app.",
    "No historical distribution is imported: no Monte Carlo, real on-time probability or risk preference in Uploaded Data Mode.",
    "Imports are immutable snapshots. Correct the source and create a new import to change inputs.",
]
FIELD_NOTES = {
    "shipment_id": "Unique trimmed shipment ID; each shipment is allocated once unless explicit unassignment is allowed.",
    "service_id": "Shared service/capacity-pool ID, declared once in capacity and referenced by options. Do not split one actual capacity into fictitious independent pools.",
    "origin": "Exact directional origin label. No geocoding or inferred equivalent locations.",
    "destination": "Exact directional destination label; must match the shared service for eligibility.",
    "weight": "Per-shipment demand, 0.001 to 1000000 kg. No conversion from pounds or tonnes.",
    "weight_unit": "Exact literal kg only.",
    "volume": "Per-shipment demand, 0.001 to 100000 m3. No conversion from litres/cubic feet.",
    "volume_unit": "Exact literal m3 only.",
    "delivery_deadline": "Full ISO instant with seconds and Z/offset, at or after common planning_start and within 87600 hours. Date-only/naive timestamps rejected.",
    "quoted_cost": "Fixed total per-shipment quoted transport cost, 0 to 1000000000 USD; decimal point, no commas/exponents. Include agreed surcharges yourself; no invented fees.",
    "currency": "Exact literal USD only. No FX conversion.",
    "expected_transit_hours": "User-provided expected elapsed hours from planning_start; at least 1/60 hour, at most87600. All options using a service_id must agree; no implied arrival guarantee.",
    "time_unit": "Exact literal h only.",
    "quote_status": "estimate or user_verified. The latter is an uploader declaration, not app verification.",
    "quote_reference": "Nonblank source/assumption reference, at most1000 characters. Required for either status.",
    "time_status": "estimate or user_verified; no reliability or variance is inferred.",
    "time_reference": "Nonblank expected-time source/assumption reference, at most1000 characters.",
    "transport_mode": "Air, Sea, Road or Rail.",
    "weight_capacity": "Total shared pool capacity in kg, from0 to1000000000; not capacity repeated per shipment.",
    "volume_capacity": "Total shared pool capacity in m3, from0 to100000000.",
    "shipment_capacity": "Total shared shipment slots; whole integer0 to1000000. Zero is valid and may make the plan infeasible.",
    "capacity_status": "estimate or user_verified; app does not confirm available carrier space.",
    "capacity_reference": "Nonblank shared-capacity source/assumption reference, at most1000 characters.",
    "planning_start": "Common readiness/dispatch instant for ALL shipments. Full ISO timestamp with Z/offset; no individual release times or service schedules in this MVP.",
    "budget": "Hard batch transport budget in USD, from0 to1000000000000; excludes virtual unassignment penalties.",
    "allow_unassigned": "Exact true or false. false requires every shipment; true adds explicit objective penalty for unassigned demand.",
    "unassigned_penalty": "Virtual USD objective penalty per unassigned shipment, from0 to1000000; not an actual invoice or financial loss.",
    "penalty_currency": "Exact USD only for objective scaling; no FX conversion.",
    "cost_weight": "Cost preference0 to100; cost+time must be positive. Never overrides hard constraints.",
    "time_weight": "Time preference0 to100; no risk preference without a historical distribution.",
    "time_limit_seconds": "Per-strategy search limit0.1 to10 seconds; the shared API worker additionally has45-second wall time.",
    "source_label": "Nonblank input source description, at most200 characters; retained with results.",
    "data_kind": "synthetic_test or user_provided. Templates must remain labeled synthetic_test unless replaced with genuine user input; provenance is not independently verified.",
}


def read_tables(files):
    """One four-sheet XLSX OR exactly four named CSV parts, bounded before parse."""
    if sum(len(content) for _, content in files.values()) > MAX_FILE_BYTES:
        raise ValueError("Planning files exceed the shared 20 MiB request budget.")
    def frame(rows, name):
        header = rows[0] if rows else []
        _validate_named_headers(header)
        names = [str(value).strip().lower() if value is not None else "" for value in header]
        if len(rows) - 1 > LIMITS[name]:
            raise ValueError(f"{name}: use at most {LIMITS[name]} rows.")
        clean = []
        for index, row in enumerate(rows[1:], 2):
            if len(row) > len(names):
                raise ValueError(f"{name} row {index}: extra fields cannot be aligned safely.")
            clean.append(["" if value is None else ("true" if value else "false") if isinstance(value, bool) else str(value) for value in row] + [""] * (len(names) - len(row)))
        return pd.DataFrame(clean, columns=names)
    if set(files) == {"workbook"}:
        filename, content = files["workbook"]
        if Path(filename).suffix.lower() != ".xlsx":
            raise ValueError("Workbook must be .xlsx.")
        validate_xlsx_archive(content)
        from openpyxl import load_workbook
        with BytesIO(content) as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                if set(book.sheetnames) != set(HEADERS):
                    raise ValueError("Workbook must contain exactly shipments, transport_options, capacity, constraints sheets.")
                tables = {}
                for name in HEADERS:
                    sheet = book[name]
                    if sheet.max_column and sheet.max_column > 128:
                        raise ValueError("Use at most 128 columns.")
                    if sheet.max_row and sheet.max_row > LIMITS[name] + 1:
                        raise ValueError(f"{name}: worksheet extends beyond {LIMITS[name]} data rows. Remove extra rows/formatting; later records cannot be silently ignored.")
                    rows = []
                    for cells in sheet.iter_rows(max_row=LIMITS[name] + 2, values_only=False):
                        if any(cell.data_type == "f" for cell in cells):
                            raise ValueError(f"{name}: formulas are not accepted; paste explicit values.")
                        row = [cell.value for cell in cells]
                        if any(value is not None and str(value).strip() for value in row):
                            rows.append(row)
                    tables[name] = frame(rows, name)
                return tables
            finally:
                book.close()
    if set(files) != set(HEADERS):
        raise ValueError("Upload one workbook OR four CSV parts named shipments, transport_options, capacity, constraints.")
    tables = {}
    for name, (filename, content) in files.items():
        if Path(filename).suffix.lower() != ".csv":
            raise ValueError(f"{name}: choose UTF-8 .csv.")
        rows = []
        for row in csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True):
            if row and any(value.strip() for value in row):
                rows.append(row)
            if len(rows) > LIMITS[name] + 1:
                raise ValueError(f"{name}: use at most {LIMITS[name]} rows.")
        tables[name] = frame(rows, name)
    return tables


def validate_tables(tables):
    errors, warnings = [], []
    error_count = 0
    def error(table, row, field, message):
        nonlocal error_count
        error_count += 1
        if len(errors) < 200:
            errors.append(dict(table=table, row=row, field=field, message=message))
    def value(row, field):
        return str(row.get(field, ""))
    def text(row, field, name, index, maximum=150):
        raw = value(row, field)
        if not raw.strip() or raw != raw.strip() or len(raw) > maximum:
            error(name, index, field, f"Use nonblank trimmed text, at most {maximum} characters.")
        return raw
    def numeric(row, field, name, index, low=0, high=1e9, integer=False):
        raw = value(row, field)
        try:
            if not re.fullmatch(r"(?:0|[1-9]\d*)(?:\.\d+)?", raw):
                raise ValueError()
            result = Decimal(raw)
            if not result.is_finite() or not Decimal(str(low)) <= result <= Decimal(str(high)) or (integer and result != result.to_integral_value()):
                raise ValueError()
            return int(result) if integer else float(result)
        except (ValueError, InvalidOperation):
            error(name, index, field, f"Use a finite {'integer' if integer else 'decimal'} between {low} and {high}; no commas, signs or exponents.")
            return 0
    def literal(row, field, choices, name, index):
        raw = value(row, field)
        if raw not in choices:
            error(name, index, field, "Accepted values: " + ", ".join(choices) + ". No automatic conversion.")
        return raw
    def instant(row, field, name, index):
        raw = value(row, field)
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:Z|[+-]\d{2}:\d{2})", raw):
                raise ValueError()
            result = datetime.fromisoformat(raw.replace("Z", "+00:00"))
            if result.tzinfo is None or result.year < 2000:
                raise ValueError()
            return result.astimezone(timezone.utc)
        except ValueError:
            error(name, index, field, "Use ISO YYYY-MM-DDTHH:MM:SSZ or an explicit offset; no naive/date-only values.")
            return None
    for name, fields in HEADERS.items():
        source = tables.get(name, pd.DataFrame())
        missing, extra = sorted(set(fields) - set(source.columns)), sorted(set(source.columns) - set(fields))
        if missing or extra:
            error(name, 1, "headers", f"Missing columns: {missing}; unexpected columns: {extra}.")
        if len(source) > LIMITS[name]:
            error(name, 1, "rows", f"At most {LIMITS[name]} rows allowed.")
    if error_count:
        return dict(valid=False, errors=errors, error_count=error_count, warnings=warnings, payload=None)
    if len(tables["constraints"]) != 1:
        error("constraints", 2, "rows", "Exactly one constraints row is required.")
    if tables["shipments"].empty:
        error("shipments", 2, "rows", "At least one shipment is required.")
    if error_count:
        return dict(valid=False, errors=errors, error_count=error_count, warnings=warnings, payload=None)
    cfg = tables["constraints"].iloc[0].to_dict()
    start = instant(cfg, "planning_start", "constraints", 2)
    budget = numeric(cfg, "budget", "constraints", 2, high=1e12)
    literal(cfg, "currency", ["USD"], "constraints", 2)
    literal(cfg, "penalty_currency", ["USD"], "constraints", 2)
    allow = literal(cfg, "allow_unassigned", ["true", "false"], "constraints", 2) == "true"
    penalty = numeric(cfg, "unassigned_penalty", "constraints", 2, high=1e6)
    cost = numeric(cfg, "cost_weight", "constraints", 2, high=100)
    duration = numeric(cfg, "time_weight", "constraints", 2, high=100)
    if cost + duration <= 0:
        error("constraints", 2, "cost_weight,time_weight", "At least one optimization preference must be positive.")
    timeout = numeric(cfg, "time_limit_seconds", "constraints", 2, low=.1, high=10)
    source_label = text(cfg, "source_label", "constraints", 2, 200)
    kind = literal(cfg, "data_kind", ["synthetic_test", "user_provided"], "constraints", 2)
    orders, service_rows, candidates, transit = [], {}, [], {}
    seen_orders, seen_pairs = set(), set()
    verified = {"quote": {"estimate": 0, "user_verified": 0}, "time": {"estimate": 0, "user_verified": 0}, "capacity": {"estimate": 0, "user_verified": 0}}
    def declaration(row, prefix, name, index):
        status = literal(row, prefix + "_status", ["estimate", "user_verified"], name, index)
        text(row, prefix + "_reference", name, index, 1000)
        if status in verified[prefix]:
            verified[prefix][status] += 1
    for index, row in enumerate(tables["shipments"].to_dict("records"), 2):
        name = "shipments"
        sid = text(row, "shipment_id", name, index)
        if sid in seen_orders: error(name, index, "shipment_id", "Duplicate shipment ID; resolve the conflict before importing.")
        seen_orders.add(sid)
        origin, dest = text(row, "origin", name, index), text(row, "destination", name, index)
        weight, volume = numeric(row, "weight", name, index, .001, 1e6), numeric(row, "volume", name, index, .001, 1e5)
        literal(row, "weight_unit", ["kg"], name, index); literal(row, "volume_unit", ["m3"], name, index)
        deadline = instant(row, "delivery_deadline", name, index)
        hours = (deadline-start).total_seconds()/3600 if start and deadline else 0
        if not 0 <= hours <= 87600: error(name, index, "delivery_deadline", "Deadline must be at/after planning_start and within 87600 hours.")
        orders.append(dict(shipment_id=sid, origin=origin, destination=dest, weight_kg=weight, volume_m3=volume, deadline_hours=hours))
    for index, row in enumerate(tables["capacity"].to_dict("records"), 2):
        name = "capacity"
        sid = text(row, "service_id", name, index)
        if sid in service_rows: error(name, index, "service_id", "Duplicate shared service ID; do not repeat its capacity per shipment.")
        for field in ("origin", "destination", "transport_mode"): text(row, field, name, index)
        literal(row, "transport_mode", ["Air", "Sea", "Road", "Rail"], name, index)
        literal(row, "weight_unit", ["kg"], name, index); literal(row, "volume_unit", ["m3"], name, index)
        weight = numeric(row, "weight_capacity", name, index)
        volume = numeric(row, "volume_capacity", name, index, high=1e8)
        slots = numeric(row, "shipment_capacity", name, index, high=1e6, integer=True)
        declaration(row, "capacity", name, index)
        service_rows[sid] = dict(service_id=sid, origin=row["origin"], destination=row["destination"], mode=row["transport_mode"], weight_capacity_kg=weight, volume_capacity_m3=volume, shipment_capacity=slots)
    for index, row in enumerate(tables["transport_options"].to_dict("records"), 2):
        name = "transport_options"
        oid, sid = text(row, "shipment_id", name, index), text(row, "service_id", name, index)
        if oid not in seen_orders: error(name, index, "shipment_id", "Unknown shipment reference.")
        if sid not in service_rows: error(name, index, "service_id", "Unknown capacity/service reference.")
        if (oid, sid) in seen_pairs: error(name, index, "shipment_id,service_id", "Duplicate shipment/service candidate.")
        seen_pairs.add((oid, sid))
        quoted = numeric(row, "quoted_cost", name, index)
        hours = numeric(row, "expected_transit_hours", name, index, 1/60, 87600)
        literal(row, "currency", ["USD"], name, index); literal(row, "time_unit", ["h"], name, index)
        if sid in transit and transit[sid] != hours: error(name, index, "expected_transit_hours", "One shared service must have one common expected transit duration. Use distinct real capacity pools for genuinely different services.")
        transit[sid] = hours
        declaration(row, "quote", name, index); declaration(row, "time", name, index)
        candidates.append(dict(shipment_id=oid, service_id=sid, cost_usd=quoted))
    unused = set(service_rows) - set(transit)
    if unused: warnings.append("Unused capacity rows are retained in source records but omitted from the solver: " + ", ".join(sorted(unused)))
    missing = seen_orders - {row["shipment_id"] for row in candidates}
    if missing: warnings.append(f"{len(missing)} shipments have no quoted candidate; strict assignment will be infeasible, soft assignment may leave them unassigned.")
    payload = None
    if not error_count:
        services = [Service(**row, mean_hours=transit[sid], std_hours=0) for sid, row in service_rows.items() if sid in transit]
        data = PlanningData(tuple(Order(**row) for row in orders), tuple(services), tuple(Candidate(**row) for row in candidates))
        payload = dict(schema=SCHEMA_VERSION, data=asdict(data), constraints=asdict(Constraints(budget, allow, penalty, timeout)),
                       preferences=asdict(Preferences(cost, duration, 0)), planning_start=start.isoformat(), source_label=source_label,
                       data_kind=kind, declarations=verified, policy=POLICY,
                       reliability="Not estimated. No historical distribution provided; internal zero std is a deterministic solver placeholder, not measured reliability.")
    return dict(valid=not error_count, errors=errors, error_count=error_count, warnings=warnings, payload=payload)
